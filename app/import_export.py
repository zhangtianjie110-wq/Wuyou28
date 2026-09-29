from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .database import Database


HEADER_ALIASES = {
    "期号": "issue_no",
    "期數": "issue_no",
    "issue": "issue_no",
    "issue_no": "issue_no",
    "预测类型": "source_type",
    "類型": "source_type",
    "类型": "source_type",
    "source": "source_type",
    "source_type": "source_type",
    "大单": "big_single",
    "大單": "big_single",
    "big_single": "big_single",
    "大双": "big_double",
    "大雙": "big_double",
    "big_double": "big_double",
    "小单": "small_single",
    "小單": "small_single",
    "small_single": "small_single",
    "小双": "small_double",
    "小雙": "small_double",
    "small_double": "small_double",
    "预测内容": "plan_text",
    "计划内容": "plan_text",
    "plan_text": "plan_text",
    "实际结果": "actual_result",
    "开奖结果": "actual_result",
    "result": "actual_result",
    "actual_result": "actual_result",
    "标记无效": "is_invalid",
    "is_invalid": "is_invalid",
}

HISTORY_HEADERS = [
    ("期号", "issue_no"),
    ("预测类型", "source_type"),
    ("大单", "big_single"),
    ("大双", "big_double"),
    ("小单", "small_single"),
    ("小双", "small_double"),
    ("预测内容", "plan_text"),
    ("实际结果", "actual_result"),
    ("结果组合", "actual_combo"),
    ("正确计划数", "correct_count"),
    ("错误计划数", "wrong_count"),
    ("最低值", "stat_min"),
    ("最高值", "stat_max"),
    ("最低两项", "lowest_two"),
    ("最低两项差值", "lowest_two_diff"),
    ("并列最低", "tied_min"),
    ("四项平均值", "average"),
    ("数据状态", "status"),
    ("异常说明", "invalid_reason"),
    ("更新时间", "updated_at"),
]


def _clean_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "")


def _mapped_row(row: dict[str, Any]) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for header, value in row.items():
        key = HEADER_ALIASES.get(_clean_header(header))
        if key:
            result[key] = "" if value is None else value
    for field in ("big_single", "big_double", "small_single", "small_double"):
        if field in result and str(result[field]).strip() != "":
            result[field] = int(float(result[field]))
    if "is_invalid" in result:
        result["is_invalid"] = str(result["is_invalid"]).strip().lower() in (
            "1",
            "true",
            "是",
            "yes",
            "无效",
        )
    if "issue_no" in result:
        value = result["issue_no"]
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        result["issue_no"] = str(value).strip()
    if "source_type" in result:
        source = str(result["source_type"]).strip().upper()
        result["source_type"] = "VIP" if source == "VIP" else str(result["source_type"]).strip()
    return result


def read_import_file(path: str | Path) -> list[dict[str, Any]]:
    file_path = Path(path)
    suffix = file_path.suffix.lower()
    if suffix == ".csv":
        last_error: Exception | None = None
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                with file_path.open("r", encoding=encoding, newline="") as handle:
                    return [_mapped_row(dict(row)) for row in csv.DictReader(handle)]
            except UnicodeDecodeError as exc:
                last_error = exc
        raise ValueError("CSV 编码无法识别，请保存为 UTF-8 或 GB18030") from last_error
    if suffix in (".xlsx", ".xlsm"):
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator, [])]
        rows = []
        for values in iterator:
            if not any(value is not None and str(value).strip() for value in values):
                continue
            rows.append(_mapped_row(dict(zip(headers, values))))
        workbook.close()
        return rows
    raise ValueError("仅支持 .xlsx、.xlsm 或 .csv 文件")


def import_file(database: Database, path: str | Path) -> dict[str, Any]:
    rows = read_import_file(path)
    summary: dict[str, Any] = {"added": 0, "updated": 0, "skipped": 0, "errors": []}
    for index, payload in enumerate(rows, start=2):
        try:
            issue = str(payload.get("issue_no") or "").strip()
            if not issue:
                raise ValueError("缺少期号")
            source = str(payload.get("source_type") or "").strip()
            if source != "VIP":
                # 允许只含“期号 + 开奖结果”的文件匹配现有 VIP 记录。
                if payload.get("actual_result") not in (None, ""):
                    matched = database.find_by_issue(issue)
                    if not matched:
                        raise ValueError("只有开奖结果但数据库中没有对应期号")
                    for existing in matched:
                        database.update_prediction(
                            existing["id"], {"actual_result": payload["actual_result"]}
                        )
                        summary["updated"] += 1
                    continue
                raise ValueError("预测类型必须是 VIP")

            existing = database.find_prediction(source, issue)
            if existing:
                updates = {key: value for key, value in payload.items() if value not in (None, "")}
                database.update_prediction(existing["id"], updates)
                summary["updated"] += 1
            else:
                database.add_prediction(payload)
                summary["added"] += 1
        except Exception as exc:
            summary["skipped"] += 1
            summary["errors"].append(f"第 {index} 行：{exc}")
    return summary


def _write_csv(path: Path, headers: list[str], rows: Iterable[list[Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def _write_xlsx(path: Path, headers: list[str], rows: Iterable[list[Any]], title: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="E9F2FF")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="24527A")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        values = [len(str(cell.value or "")) for cell in column[:200]]
        sheet.column_dimensions[column[0].column_letter].width = min(max(values + [8]) + 2, 35)
    workbook.save(path)


def export_history(path: str | Path, records: list[dict[str, Any]]) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    headers = [header for header, _ in HISTORY_HEADERS]
    rows = [[record.get(field, "") for _, field in HISTORY_HEADERS] for record in records]
    if destination.suffix.lower() == ".csv":
        _write_csv(destination, headers, rows)
    elif destination.suffix.lower() == ".xlsx":
        _write_xlsx(destination, headers, rows, "历史数据")
    else:
        raise ValueError("导出文件必须是 .xlsx 或 .csv")
    return destination


def export_backtest(path: str | Path, strategy: dict[str, Any], result: dict[str, Any]) -> Path:
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    headers = [
        "策略名称",
        "预测类型",
        "期号",
        "数据集",
        "选择组合",
        "实际组合",
        "是否命中",
        "最低值",
        "最低两项差值",
        "上一期错误数",
        "进入前连续错",
    ]
    rows = [
        [
            strategy["name"],
            strategy["source_type"],
            item["issue_no"],
            item["dataset"],
            item["selected"],
            item["actual_combo"],
            "命中" if item["hit"] else "未命中",
            item["min_value"],
            item["lowest_two_diff"],
            "" if item["previous_wrong"] is None else item["previous_wrong"],
            item["raw_loss_streak_before"],
        ]
        for item in result.get("details", [])
    ]
    if destination.suffix.lower() == ".csv":
        _write_csv(destination, headers, rows)
    elif destination.suffix.lower() == ".xlsx":
        _write_xlsx(destination, headers, rows, "回测明细")
    else:
        raise ValueError("导出文件必须是 .xlsx 或 .csv")
    return destination


def create_import_template(path: str | Path) -> Path:
    destination = Path(path)
    headers = ["期号", "预测类型", "大单", "大双", "小单", "小双", "预测内容", "实际结果"]
    example = ["20260923001", "VIP", 25, 25, 25, 25, "", "大单"]
    if destination.suffix.lower() == ".csv":
        _write_csv(destination, headers, [example])
    elif destination.suffix.lower() == ".xlsx":
        _write_xlsx(destination, headers, [example], "导入模板")
    else:
        raise ValueError("模板文件必须是 .xlsx 或 .csv")
    return destination
